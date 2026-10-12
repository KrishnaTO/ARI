"""Compare a master list's ALPHABETICAL sheet with the active diseases in the ARI ontology,
column by column, and write the differences to an Excel report.

Usage: python compare.py <master.xlsx> <ari_t1d.owl> <out.xlsx>
"""
import re, sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from sheet_items import items, norm, same_item

MASTER, OWL_PATH, OUT = sys.argv[1:4]

RDF = '{http://www.w3.org/1999/02/22-rdf-syntax-ns#}'
RDFS = '{http://www.w3.org/2000/01/rdf-schema#}'
A = '{http://www.aurint.org/ontologies/ari-disease#}'

# (report column, sheet header prefix, ARI property or properties, kind)
COLUMNS = [
    ('Veredict', 'Veredict', 'diseaseCategory', 'value'),
    ('System', 'System', 'targetsTissue', 'system'),
    ('Target Antigens', 'Target Antigens', 'involvesAntigen', 'list'),
    ('Target Cells', 'Target Cells', 'affectsCellTissueOrOrgan', 'list'),
    ('Antibodies', 'Antibodies', 'hasAntibody', 'list'),
    ('Self Reactive Cells', 'Self Reactive Cells', ('involvesTCell', 'involvesAPC'), 'list'),
    ('Symptoms', 'Symptoms', 'hasSymptom', 'list'),
    ('Molecular Mimicry', 'Molecular Mimicry', 'hasEtiologyFactor', 'list'),
    ('Environmental Triggers', 'Enviromental Triggers', 'hasEnvironmentalTrigger', 'list'),
    ('Treatments', 'Treatments', 'hasTreatment', 'list'),
]
ARI_FIELD_TEXT = {
    'Veredict': 'diseaseCategory',
    'System': 'targetsTissue (tissue labels)',
    'Target Antigens': 'involvesAntigen',
    'Target Cells': 'affectsCellTissueOrOrgan',
    'Antibodies': 'hasAntibody',
    'Self Reactive Cells': 'involvesTCell + involvesAPC',
    'Symptoms': 'hasSymptom',
    'Molecular Mimicry': 'hasEtiologyFactor "Molecular mimicry" (description)',
    'Environmental Triggers': 'hasEnvironmentalTrigger',
    'Treatments': 'hasTreatment',
}
# Sheet names that differ from every ARI label and synonym but name the same disease.
MANUAL = {
    'Anti-DPX encephalitis': ('ARI:0001122', 'DPPX is dipeptidyl-peptidase-like protein 6 (DPP6). '
                              'ARI:0001121 Limbic encephalitis also has "Anti-DPPX encephalitis" as its definition.'),
    'Anti-GABA A encephalitis': ('ARI:0001082', ''),
    'Anti-IgLON5 disease': ('ARI:0001083', ''),
    'Anti-LGI1 autoimmune encephalitis': ('ARI:0001123', ''),
    'Autoimmune autonomic ganglionopathy': ('ARI:0001030', ''),
    'Immune dysregulation, polyendocrinopathy, enteropathy, X-linked': ('ARI:0001106', ''),
    'Neonatal lupus / congenital heart block due to maternal anti-Ro/SSA and anti-La/SSB': ('ARI:0001141', ''),
    'Primary membranous nephropathy (PLA2R/THSD7A-associated)': ('ARI:0001102', ''),
    'Autoimmune postural orthostatic tachycardia syndrome (POTS)': ('ARI:0001164', 'ARI label has no "autoimmune".'),
    'Autoimmune small fiber neuropathy': ('ARI:0001191', 'ARI label has no "autoimmune".'),
    'Autoimmune optic neuritis (non-MS, non-NMOSD)': ('ARI:0001042', 'Sheet says neuritis, ARI says neuropathy.'),
    'Autoimmune chylomicronemia due to anti-GPIHBP1 antibodies': ('ARI:0001165', 'ARI label is the SNOMED term.'),
    'Anti-NXP2 dermatomyositis (juvenile form with calcinosis)': ('ARI:0001059', 'ARI label is the SNOMED term; '
                                                                  'it does not name NXP2.'),
    'Autoinflammatory calcifying aortic valve disease (immune-mediated)': ('ARI:0001020', 'ARI label is the SNOMED term.'),
}
# ARI-only diseases that are not a child of a disease in the sheet.
ARI_ONLY = {
    'ARI:0001098': 'Hemophilia B Leyden is inherited. The sheet\'s autoimmune form is ARI:0001223 '
                   'Autoimmune hemophilia B (anti-factor IX antibody).',
    'ARI:0001121': 'Its definition reads "Anti-DPPX encephalitis" and its definition source names "Anti-AMPAR '
                   'encephalitis". The sheet\'s Anti-DPX encephalitis matched ARI:0001122 and its Anti-AMPAR '
                   'encephalitis is ARI:0001217.',
    'ARI:0001145': 'Subtype of Myasthenia gravis (ARI:0001136), which is in the sheet; ARI has no parent link for it.',
    'ARI:0001183': 'Rheumatoid aortitis is aortitis in rheumatoid arthritis. The sheet\'s isolated aortitis is '
                   'ARI:0001219 Autoimmune aortitis (isolated, non-vasculitic).',
}

STATUS_FILL = {
    'Same': 'E2EFDA', 'Partly same': 'FFF2CC', 'Changed': 'FCE4D6', 'Differs': 'FCE4D6',
    'New in sheet': 'DDEBF7', 'Only in ARI': 'FFF2CC', 'Disease not in ARI': 'EDEDED',
    'Unclear in sheet': 'F8CBAD',
}


# ---------- ARI ----------
root = ET.parse(OWL_PATH).getroot()
# Top-level nodes, whether written as owl:NamedIndividual or as a typed element such as <Symptom>.
nodes = {e.get(RDF + 'about'): e for e in root if e.get(RDF + 'about')}
ari = {}
for e in nodes.values():
    i = e.findtext(A + 'ARI_ID')
    if i and i.startswith('ARI:') and e.findtext(A + 'ARI_Obsolete') != 'true':
        ari[i] = e
label = {i: e.findtext(RDFS + 'label') for i, e in ari.items()}
id_by_iri = {iri: e.findtext(A + 'ARI_ID') for iri, e in nodes.items() if e.findtext(A + 'ARI_ID')}


def linked(e, prop):
    return [nodes[r.get(RDF + 'resource')] for r in e.findall(A + prop)]


def ari_items(i, prop):
    e = ari[i]
    if isinstance(prop, tuple):
        return [v for p in prop for v in ari_items(i, p)]
    if prop == 'diseaseCategory':
        v = e.findtext(A + prop)
        return [v] if v else []
    if prop == 'hasEtiologyFactor':
        return [n.findtext(A + 'etiologyDescription') for n in linked(e, prop)
                if n.findtext(RDFS + 'label') == 'Molecular mimicry']
    return [n.findtext(RDFS + 'label') for n in linked(e, prop)]


names = defaultdict(set)
for i, e in ari.items():
    names[norm(label[i])].add((i, 'Name = ARI label'))
    for s in e.findall(A + 'ARI_Synonym'):
        names[norm(s.text)].add((i, 'Name = ARI synonym'))

# ---------- master list ----------
ws = openpyxl.load_workbook(MASTER, data_only=True)['ALPHABETICAL']
headers = [str(c.value or '').strip() for c in ws[1]]


def col(prefix):
    hits = [k for k, h in enumerate(headers) if h.startswith(prefix)]
    if len(hits) != 1:
        raise SystemExit(f'Expected one column starting {prefix!r}, found {len(hits)}')
    return hits[0]


NAME, MESH = col('Disease'), col('Mesh Terms')
COLIDX = {c[0]: col(c[1]) for c in COLUMNS}
rows = [(n, r) for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2) if r[NAME]]

for s, (i, _) in MANUAL.items():
    if i not in ari or s not in {str(r[NAME]).strip() for _, r in rows}:
        raise SystemExit(f'Manual match {s!r} -> {i} is stale')


def lookup(cands):
    hits = set()
    for c in cands:
        hits |= names.get(norm(c), set())
    ids = {h[0] for h in hits}
    if len(ids) > 1:
        raise SystemExit(f'{cands} matches several ARI diseases: {sorted(ids)}')
    return hits


def match(name, mesh):
    hits = lookup([name])
    if hits:
        return next(iter(hits))[0], sorted(h[1] for h in hits)[0], ''
    variants = [re.sub(r'\([^)]*\)', '', name)] + re.findall(r'\(([^)]*)\)', name)
    variants += [p for v in variants + [name] for p in v.split('/')]
    hits = lookup([v.strip() for v in variants if v.strip()])
    if hits:
        return next(iter(hits))[0], 'Name variant', 'Matched after dropping the parenthetical or splitting on "/".'
    terms = [t.strip() for t in (mesh or '').split('\n') if t.strip()]
    hits = lookup(terms)
    if hits:
        return next(iter(hits))[0], 'MeSH term', 'Sheet name differs; one of its MeSH terms is an ARI name.'
    if name in MANUAL:
        i, note = MANUAL[name]
        return i, 'Manual', note
    return None, 'Not in ARI', ''


matched = []   # (sheet row, sheet name, ARI ID or None, method, note, row values)
seen = {}
for n, r in rows:
    name = str(r[NAME]).strip()
    i, how, note = match(name, r[MESH])
    if i and i in seen:
        first = seen[i]
        same = all(r[k] == first[5][k] for k in COLIDX.values())
        note = (f'Duplicate of sheet row {first[0]}; the ten compared columns '
                + ('are identical.' if same else 'differ - compare both rows.'))
        matched.append((n, name, i, 'Duplicate row', note, r))
        continue
    entry = (n, name, i, how, note, r)
    matched.append(entry)
    if i:
        seen[i] = entry

parents_of = defaultdict(list)
for i, e in ari.items():
    for p in e.findall(A + 'hasParentDisease'):
        parents_of[i].append(id_by_iri[p.get(RDF + 'resource')])
ari_only = sorted(set(ari) - set(seen))


# ---------- comparison ----------
def system_key(s):
    return re.sub(r' system$', '', norm(s))


system_vocab = {system_key(t) for i in ari for t in ari_items(i, 'targetsTissue')}
system_vocab |= {system_key(r[COLIDX['System']]) for _, r in rows
                 if r[COLIDX['System']] and '\n' not in str(r[COLIDX['System']]).strip()}


def compare(column, kind, prop, i, raw):
    """Return (status, sheet value, ARI value, only in sheet, only in ARI, note)."""
    sheet_txt = str(raw).strip() if raw is not None else ''
    if i is None:
        return 'Disease not in ARI', sheet_txt, '', '', '', ''
    ari_list = ari_items(i, prop)
    ari_txt = '\n'.join(ari_list)
    if kind == 'value':
        if not ari_list:
            return ('New in sheet' if sheet_txt else 'No data'), sheet_txt, '', '', '', ''
        return ('Same' if norm(sheet_txt) == norm(ari_list[0]) else 'Changed'), sheet_txt, ari_txt, '', '', ''
    if kind == 'system':
        lines = [l.strip().split(' - ')[0].strip() for l in sheet_txt.split('\n') if l.strip()]
        systems = {system_key(l) for l in lines if system_key(l) in system_vocab}
        note = 'System cell also holds free text; needs cleanup.' if len(systems) < len(lines) else ''
        theirs = {system_key(t) for t in ari_list}
        if not systems:
            status = 'Unclear in sheet'
        elif not theirs:
            status = 'New in sheet'
        elif systems == theirs:
            status = 'Same'
        elif systems & theirs:
            status = 'Partly same'
        else:
            status = 'Changed'
        return status, sheet_txt, ari_txt, '', '', note
    mine = items(raw)
    if not mine and not ari_list:
        return 'No data', sheet_txt, '', '', '', ''
    if not ari_list:
        return 'New in sheet', sheet_txt, '', '', '', ''
    if not mine:
        return 'Only in ARI', sheet_txt, ari_txt, '', '', ''
    only_s = [m for m in mine if not any(same_item(m, a) for a in ari_list)]
    only_a = [a for a in ari_list if not any(same_item(m, a) for m in mine)]
    status = 'Same' if not only_s and not only_a else 'Differs'
    return status, sheet_txt, ari_txt, '\n'.join(only_s), '\n'.join(only_a), ''


results = {c[0]: [] for c in COLUMNS}
for n, name, i, how, note, r in matched:
    if how == 'Duplicate row':
        continue
    for column, _, prop, kind in COLUMNS:
        results[column].append((n, name, i, label.get(i, '')) + compare(column, kind, prop, i, r[COLIDX[column]]))


# ---------- report ----------
FONT, BOLD = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True)
HEAD_FILL = PatternFill('solid', fgColor='D9E1F2')
WRAP = Alignment(wrap_text=True, vertical='top')


def table(ws, header, data, widths, status_col=None):
    ws.append(header)
    for row in data:
        ws.append(list(row))
    for cells in ws.iter_rows():
        for c in cells:
            c.font, c.alignment = FONT, WRAP
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD_FILL
    if status_col is not None:
        for cells in ws.iter_rows(min_row=2):
            fill = STATUS_FILL.get(cells[status_col].value)
            if fill:
                cells[status_col].fill = PatternFill('solid', fgColor=fill)
    for k, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(k + 1)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


wb = openpyxl.Workbook()
summary = wb.active
summary.title = 'Summary'

match_rows = [(n, name, i or '', label.get(i, ''), how, note) for n, name, i, how, note, _ in matched]
for i in ari_only:
    note = ARI_ONLY.get(i) or '; '.join(f'Child of {p} {label[p]}, which is in the sheet as "{seen[p][1]}".'
                                        for p in parents_of[i] if p in seen)
    match_rows.append(('', '', i, label[i], 'Not in sheet', note))
table(wb.create_sheet('Disease Match'), ['Sheet row', 'Sheet disease', 'ARI ID', 'ARI label', 'Match', 'Note'],
      match_rows, [10, 45, 13, 45, 20, 70])

for column, _, _, kind in COLUMNS:
    ws = wb.create_sheet(column)
    if kind == 'list':
        header = ['Sheet row', 'Sheet disease', 'ARI ID', 'ARI label', 'Status',
                  'Only in sheet', 'Only in ARI', f'Sheet: {column}', f'ARI: {ARI_FIELD_TEXT[column]}']
        data = [r[:5] + r[7:9] + r[5:7] for r in results[column]]
        table(ws, header, data, [10, 40, 13, 40, 18, 45, 45, 60, 45], status_col=4)
    else:
        header = ['Sheet row', 'Sheet disease', 'ARI ID', 'ARI label', 'Status',
                  f'Sheet: {column}', f'ARI: {ARI_FIELD_TEXT[column]}', 'Note']
        data = [r[:7] + r[9:10] for r in results[column]]
        table(ws, header, data, [10, 40, 13, 40, 18, 40, 30, 45], status_col=4)

statuses = ['Same', 'Partly same', 'Changed', 'Differs', 'New in sheet', 'Only in ARI', 'No data',
            'Unclear in sheet', 'Disease not in ARI']
how_count = Counter(m[3] for m in matched)
lines = [
    ['ARI master list vs ARI ontology - column comparison'],
    [f'Sheet: {MASTER.replace(chr(92), "/").split("/")[-1]} (ALPHABETICAL)'],
    [f'ARI: ontologies/ari_t1d.owl, active diseases only (ARI_Obsolete true excluded). Generated {date.today()}.'],
    [],
    ['Diseases', 'Count'],
    ['Sheet rows', len(rows)],
    ['Sheet rows that repeat a disease', how_count['Duplicate row']],
    ['Active ARI diseases', len(ari)],
    ['Matched to an ARI disease', len(seen)],
] + [[f'   matched by: {h}', how_count[h]] for h in
     ['Name = ARI label', 'Name = ARI synonym', 'Name variant', 'MeSH term', 'Manual']] + [
    ['In the sheet, not in ARI', how_count['Not in ARI']],
    ['In ARI, not in the sheet', len(ari_only)],
    [],
    ['Column', 'ARI field'] + statuses,
] + [[c[0], ARI_FIELD_TEXT[c[0]]] + [Counter(r[4] for r in results[c[0]])[s] for s in statuses] for c in COLUMNS] + [
    [],
    ['How to read'],
    ['One row per disease in the sheet; a repeated disease is compared once (see Disease Match).'],
    ['Same / Changed: one value each side. Partly same: the systems overlap but are not equal.'],
    ['Same / Differs: list columns, compared item by item. Two items count as the same when more than half '
     'the words of the shorter one appear in the other, ignoring frequencies such as (80-90%).'],
    ['Only in sheet / Only in ARI list the items with no counterpart. Wording can differ for the same idea, '
     'so read them against the full values.'],
    ['New in sheet: ARI has nothing in that field. Only in ARI: the sheet cell is empty or a placeholder.'],
    ['Placeholders "Unknown", "Unconfirmed" and "Unknown / unconfirmed" count as no data, except in Veredict.'],
    ['Unclear in sheet: the System cell holds text but no recognisable system name.'],
]
for line in lines:
    summary.append(line)
for cells in summary.iter_rows():
    for c in cells:
        c.font, c.alignment = FONT, Alignment(vertical='top')
summary['A1'].font = Font(name='Arial', size=12, bold=True)
for cells in summary.iter_rows():
    if cells[0].value in ('Diseases', 'Column', 'How to read'):
        for c in cells:
            c.font = BOLD
            if c.value is not None:
                c.fill = HEAD_FILL
summary.column_dimensions['A'].width = 34
summary.column_dimensions['B'].width = 48
for k in range(len(statuses)):
    summary.column_dimensions[openpyxl.utils.get_column_letter(k + 3)].width = 13

wb.save(OUT)
print(f'{len(rows)} sheet rows, {len(seen)} matched, {how_count["Not in ARI"]} sheet-only, '
      f'{len(ari_only)} ARI-only -> {OUT}')
