"""Import the master list's prevalence and incidence into ari_t1d.owl, for every disease that
report 10 matched. The sheet has a worldwide and a US figure, each marked incidence or
prevalence:

- worldwide incidence -> incidenceWorldPer100k
- worldwide prevalence -> prevalenceWorldPer100k
- US incidence -> incidenceUSPer100k
- US prevalence -> prevalencePer100k (the editor reads it as the US rate)
- every figure, verbatim, with the sheet's notes on it -> ARI_PrevalenceDesc

A rate is stored only for a plain "1.0 cases per N" or "per A to B" (the midpoint of the range,
per 100,000). Anything else, such as "Extremely rare" or a rate per live birth, is kept as text
only. The values the disease had in these five fields are replaced, and the change-log entry
records them.

Usage: python import_prevalence.py <master.xlsx> <ari_t1d.owl> <10_Master_List_Comparison.xlsx>
"""
import html, re, sys
from datetime import datetime, timezone
from decimal import Decimal
from xml.sax.saxutils import escape

import openpyxl

MASTER, OWL_PATH, REPORT = sys.argv[1:4]
SOURCE = 'ARI Master List 2026-10-08'
CURATOR = 'Leon Islas'
XSD = 'http://www.w3.org/2001/XMLSchema#'
DISEASE_IRI = 'https://diseases.autoimmuneregistry.org/disease/ARI_'
FIELDS = ('prevalencePer100k', 'prevalenceWorldPer100k', 'incidenceWorldPer100k', 'incidenceUSPer100k',
          'ARI_PrevalenceDesc')
# (region as written in ARI_PrevalenceDesc, sheet value header prefix, field for incidence, for prevalence)
REGIONS = [('Worldwide', 'Global Incidence/Prevalence', 'incidenceWorldPer100k', 'prevalenceWorldPer100k'),
           ('US', 'Incidence/Prevalence in the US', 'incidenceUSPer100k', 'prevalencePer100k')]
RATE = re.compile(r'1(?:\.0)? cases? per ([\d,]+)(?: to ([\d,]+))?', re.I)
RARE = re.compile(r'\s*\(extremely rare\)\s*$|^extremely rare$', re.I)

ws = openpyxl.load_workbook(MASTER)['ALPHABETICAL']
headers = [str(c.value or '').strip() for c in ws[1]]
cols = {}
for region, prefix, _, _ in REGIONS:
    hits = [k + 1 for k, h in enumerate(headers) if h.startswith(prefix)]
    if len(hits) != 1 or not headers[hits[0] - 2].startswith('prevtype'):
        raise SystemExit(f'Expected one {prefix!r} column with a prevtype column before it')
    cols[region] = (hits[0] - 1, hits[0])     # (type column, value column)

rows_of = {}
for r in openpyxl.load_workbook(REPORT)['Disease Match'].iter_rows(min_row=2, values_only=True):
    if r[0] and r[2]:
        rows_of.setdefault(r[2], []).append(r[0])


def figure(row, region):
    """(incidence|prevalence, verbatim text, notes on the cells, rate per 100k or None)."""
    type_col, value_col = cols[region]
    kind = str(ws.cell(row, type_col).value or '').strip().lower()
    if kind not in ('incidence', 'prevalence'):
        raise SystemExit(f'Row {row}: {region} type is {kind!r}')
    lines = [l.strip() for l in str(ws.cell(row, value_col).value or '').split('\n') if l.strip()]
    notes = []
    for c in (type_col, value_col):
        comment = ws.cell(row, c).comment
        if comment and comment.text.strip() not in notes:
            notes.append(comment.text.strip())
    plain = [RARE.sub('', l).strip() for l in lines]
    plain = [l for l in plain if l]
    rate = None
    if len(plain) == 1 and RATE.fullmatch(plain[0]):
        m = RATE.fullmatch(plain[0])
        a = int(m.group(1).replace(',', ''))
        b = int(m.group(2).replace(',', '')) if m.group(2) else a
        rate = (100000 / a + 100000 / b) / 2
    return kind, '; '.join(lines), notes, rate


def decimal_text(x):
    return format(Decimal(repr(float(f'{x:.4g}'))), 'f')


def describe(row, region):
    kind, text, notes, _ = figure(row, region)
    out = f'{region} {kind}: {text}.' if text else ''
    for n in notes:
        out += f' Note on the {region} figure: {n}'
    return out.strip()


text = open(OWL_PATH, encoding='utf-8', newline='').read()
nl = '\r\n' if '\r\n' in text else '\n'
lines = text.split(nl)
span = {}
for k, line in enumerate(lines):
    m = re.match(r'<([\w:]+) rdf:about="' + re.escape(DISEASE_IRI) + r'(\d{7})">$', line)
    if m:
        close = next(j for j in range(k + 1, len(lines)) if lines[j] == f'</{m.group(1)}>')
        span['ARI:' + m.group(2)] = (k, close)

now = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%S+00:00')
drop, before, counts = set(), {}, {f: 0 for f in FIELDS}
for ari, rows in rows_of.items():
    row, duplicates = rows[0], rows[1:]
    open_k, close_k = span[ari]
    old = {}
    for k in range(open_k + 1, close_k):
        m = re.match(r'\s*<(' + '|'.join(FIELDS) + r') [^>]*>([^<]*)</\1>$', lines[k])
        if m:
            old.setdefault(m.group(1), []).append(html.unescape(m.group(2)))
            drop.add(k)
    new, said = {}, []
    for region, _, inc_field, prev_field in REGIONS:
        kind, txt, _, rate = figure(row, region)
        field = inc_field if kind == 'incidence' else prev_field
        if rate is not None:
            new[field] = decimal_text(rate)
            said.append(f'{field} {new[field]} from the {region} {kind} "{txt}"'
                        + (' (midpoint of the range)' if ' to ' in txt else ''))
    desc = ' '.join(d for d in (describe(row, r[0]) for r in REGIONS) if d)
    for dup in duplicates:
        for region, *_ in REGIONS:
            if figure(dup, region)[:2] != figure(row, region)[:2]:
                kind, txt = figure(dup, region)[:2]
                desc += f' Sheet row {dup}, a duplicate of row {row}, gives {region} {kind}: {txt}.'
    new['ARI_PrevalenceDesc'] = desc
    add = []
    for field in FIELDS:
        if field in new:
            dtype = 'string' if field == 'ARI_PrevalenceDesc' else 'decimal'
            add.append(f'  <{field} rdf:datatype="{XSD}{dtype}">{escape(new[field])}</{field}>')
            counts[field] += 1
    note = (f'Prevalence and incidence from {SOURCE} from {CURATOR}, ALPHABETICAL row {row}: '
            + '; '.join(said + ['ARI_PrevalenceDesc rewritten with every figure and note']) + '.')
    replaced = [f'{f} {v!r}' for f in FIELDS for v in old.get(f, []) if new.get(f) != v]
    if replaced:
        note += ' Replaced: ' + '; '.join(replaced) + '.'
    add.append(f'  <ARI_ChangeLog rdf:datatype="{XSD}string">{now} | {CURATOR} | {escape(note)}</ARI_ChangeLog>')
    before[close_k] = add

out = []
for k, line in enumerate(lines):
    out += before.get(k, [])
    if k not in drop:
        out.append(line)
open(OWL_PATH, 'w', encoding='utf-8', newline='').write(nl.join(out))
print(f'{len(before)} diseases; replaced {len(drop)} old values')
for f, n in counts.items():
    print(f'  {f}: {n}')
