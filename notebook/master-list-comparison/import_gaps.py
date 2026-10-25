"""Import the master-list data that report 12 marks "In repo - import" into ari_disease.owl, for every
disease that report 10 matched. Only fields a disease does not yet have are filled:

  Definition                   -> rdfs:comment
  Veredict                     -> diseaseCategory
  Association with other AI    -> associatedWithAutoimmune (link to the ARI disease the name matches)
                                  and ARI_AssociatedDisease (each line verbatim, matched or not)
  Female/Male ratio, Countries -> demographicBias
  PATHOPHYSIOLOGY              -> hasPathwayStep, one PathwayStep per numbered step
  Symptoms                     -> hasSymptom, one Symptom per line, frequency as likelihood
  Genes (autoinflammation)     -> hasGeneticAssociation to an HLAAssociation, for HLA alleles

Not imported, by decision: Evidence Quality (the master's grade column is for autoinflammation
evidence, unconfirmed as the data model's grade), MeSH terms as synonyms (left to synonym review)
and report 11 parents (all remaining ones were rejected in proposed_parents.tsv).

Written against the layout ari_disease.owl had before the canonical formatter (relative `#` IRIs,
one element per line). Check its lookups against today's file before re-running it, then run
tools/format-owl.sh on the result.

Usage: python import_gaps.py <master.xlsx> <ari_disease.owl> <10_Master_List_Comparison.xlsx> <ari.synonym_blocklist.tsv>
"""
import re, sys
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from xml.sax.saxutils import escape

import openpyxl

from sheet_items import items, norm

MASTER, OWL_PATH, REPORT, BLOCKLIST = sys.argv[1:5]
SOURCE = 'ARI Master List 2026-10-08'
CURATOR = 'Leon Islas'
STR = ' rdf:datatype="http://www.w3.org/2001/XMLSchema#string"'
INT = ' rdf:datatype="http://www.w3.org/2001/XMLSchema#integer"'
DISEASE_IRI = 'https://diseases.autoimmuneregistry.org/disease/ARI_'
A = '{http://www.aurint.org/ontologies/ari-disease#}'
RDF = '{http://www.w3.org/1999/02/22-rdf-syntax-ns#}'
RDFS = '{http://www.w3.org/2000/01/rdf-schema#}'
FREQ = re.compile(r'\s*(?:\(([^()]*\d[^()]*%[^()]*)\)|-\s*(\d[\d.,–-]*\s*%))\s*$')
STEP = re.compile(r'^\s*(\d+)[.)]\s*(.+)$')
HLA = re.compile(r'\bHLA\b|human leukocyte antigen', re.I)
ASSOC_PROP, ASSOC_TEXT = 'associatedWithAutoimmune', 'ARI_AssociatedDisease'

text = open(OWL_PATH, encoding='utf-8', newline='').read()
nl = '\r\n' if '\r\n' in text else '\n'
lines = text.split(nl)
root = ET.fromstring(text)
nodes = {e.get(RDF + 'about') for e in root if e.get(RDF + 'about')}
for decl in (f'#{ASSOC_PROP}', f'#{ASSOC_TEXT}'):
    if f'rdf:about="{decl}"' in text:
        raise SystemExit(f'{decl} is already declared')

diseases = {e.findtext(A + 'ARI_ID'): e for e in root
            if any(t.get(RDF + 'resource') == '#AutoimmuneDisease' for t in e.findall(RDF + 'type'))}

# Names that identify one ARI disease: labels and synonyms, less the blocklisted synonyms, which
# name a different disease than the one they are stored on.
blocked = {(r.split('\t')[0], norm(r.split('\t')[1])) for r in open(BLOCKLIST, encoding='utf-8').read().splitlines()
           if r and not r.startswith('#') and r.count('\t') >= 1}
name_to_ari = {}
for ari, e in diseases.items():
    if 'true' in [o.text for o in e.findall(A + 'ARI_Obsolete')]:
        continue
    for n in [e.findtext(RDFS + 'label')] + [s.text for s in e.findall(A + 'ARI_Synonym')]:
        if (ari, norm(n)) not in blocked:
            name_to_ari.setdefault(norm(n), set()).add(ari)


def associated(line, self_ari):
    """ARI diseases a sheet line names: the whole line, or each part of an 'A / B' line."""
    bare = re.sub(r'\s*\([^)]*\)\s*$', '', line).strip()
    for candidate in ([bare], [p.strip() for p in bare.split(' / ')]):
        hits = [name_to_ari.get(norm(c), set()) for c in candidate]
        found = [next(iter(h)) for h in hits if len(h) == 1 and h != {self_ari}]
        if found:
            return found
    return []


ws = openpyxl.load_workbook(MASTER, data_only=True)['ALPHABETICAL']
headers = [str(c.value or '').strip() for c in ws[1]]


def col(prefix, nth=0):
    hits = [k for k, h in enumerate(headers) if h.startswith(prefix)]
    if len(hits) <= nth:
        raise SystemExit(f'No column starting {prefix!r}')
    return hits[nth] + 1


C_DEF = col('Autoimune Disease Definition (according to \nAutoimmune')
C_VERDICT = col('Veredict')
C_ASSOC = col('Association with other Autoimmune')
C_FM = col('Female')
C_REGIONS = col('Countries and Regions')
C_PATHO = col('PATHOPHYSIOLOGY')
C_SYMPTOMS = col('Symptoms')
C_GENES = col('1. Genes/Genetic Regions', 1)
match = [(r[0], r[2]) for r in openpyxl.load_workbook(REPORT)['Disease Match'].iter_rows(min_row=2, values_only=True)
         if r[0] and r[2] and r[4] != 'Duplicate row']

span = {}
for k, line in enumerate(lines):
    m = re.match(r'<([\w:]+) rdf:about="' + re.escape(DISEASE_IRI) + r'(\d{7})">$', line)
    if m:
        close = next(j for j in range(k + 1, len(lines)) if lines[j] == f'</{m.group(1)}>')
        span['ARI:' + m.group(2)] = (k, close)


def has(e, prop, ns=A):
    return any((c.text or c.get(RDF + 'resource') or '').strip() for c in e.findall(ns + prop))


def raw(row, c):
    v = ws.cell(row, c).value
    return str(v).strip() if v is not None else ''


def node(iri, cls, props):
    block = [f'<owl:NamedIndividual rdf:about="{iri}">', f'  <rdf:type rdf:resource="#{cls}"/>']
    for prop, dtype, value in props:
        block.append(f'  <{prop}{dtype}>{escape(str(value))}</{prop}>')
    return block + ['</owl:NamedIndividual>']


def new_iri(prefix, ari):
    k = 1
    while f'#{prefix}_{ari[4:]}_{k}' in nodes:
        k += 1
    iri = f'#{prefix}_{ari[4:]}_{k}'
    nodes.add(iri)
    return iri


now = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%S+00:00')
after, before, blocks, totals = {}, {}, [], {}
for row, ari in match:
    e = diseases[ari]
    links, values, done = [], [], []

    definition = raw(row, C_DEF)
    if items(definition) and not has(e, "comment", RDFS):
        values.append(f'  <rdfs:comment{STR}>{escape(definition)}</rdfs:comment>')
        done.append('definition')

    verdict = raw(row, C_VERDICT)
    if verdict and not has(e, "diseaseCategory"):
        values.append(f'  <diseaseCategory{STR}>{escape(verdict)}</diseaseCategory>')
        done.append(f'disease type {verdict}')

    assoc_lines, assoc_links = items(raw(row, C_ASSOC)), []
    if assoc_lines and not has(e, ASSOC_TEXT):
        for line in assoc_lines:
            values.append(f'  <{ASSOC_TEXT}{STR}>{escape(line)}</{ASSOC_TEXT}>')
            for other in associated(line, ari):
                if other not in assoc_links:
                    assoc_links.append(other)
        links += [f'  <{ASSOC_PROP} rdf:resource="{DISEASE_IRI}{a[4:]}"/>' for a in assoc_links]
        done.append(f'{len(assoc_lines)} associated diseases ({len(assoc_links)} linked to ARI diseases)')

    bias = []
    fm = raw(row, C_FM)
    if fm and items(fm):
        bias.append(f'Female:male ratio {fm}')
    regions = items(raw(row, C_REGIONS))
    if regions:
        bias.append('Regions with very high prevalence: ' + '; '.join(regions))
    if bias and not has(e, 'demographicBias'):
        values.append(f'  <demographicBias{STR}>{escape(". ".join(bias))}</demographicBias>')
        done.append('demographic bias')

    steps = items(raw(row, C_PATHO))
    if steps and not has(e, 'hasPathwayStep'):
        for k, line in enumerate(steps, 1):
            m = STEP.match(line)
            order, description = (int(m.group(1)), m.group(2).strip()) if m else (k, line)
            iri = new_iri('Path', ari)
            blocks.append(node(iri, 'PathwayStep', [('rdfs:label', STR, f'Step {order}'), ('stepOrder', INT, order),
                                                    ('stepDescription', STR, description)]))
            links.append(f'  <hasPathwayStep rdf:resource="{iri}"/>')
        done.append(f'{len(steps)} pathway steps')

    symptoms = items(raw(row, C_SYMPTOMS))
    if symptoms and not has(e, 'hasSymptom'):
        for line in symptoms:
            freq = FREQ.search(line)
            props = [('rdfs:label', STR, line[:freq.start()].strip() if freq else line)]
            if freq:
                props.append(('likelihood', STR, (freq.group(1) or freq.group(2)).strip()))
            iri = new_iri('Sym', ari)
            blocks.append(node(iri, 'Symptom', props + [('ARI_Obsolete', STR, 'false')]))
            links.append(f'  <hasSymptom rdf:resource="{iri}"/>')
        done.append(f'{len(symptoms)} symptoms')

    alleles = [l for l in items(raw(row, C_GENES)) if HLA.search(l)]
    if alleles and not has(e, 'hasGeneticAssociation'):
        for line in alleles:
            iri = new_iri('HLA', ari)
            blocks.append(node(iri, 'HLAAssociation', [('rdfs:label', STR, line), ('ARI_Obsolete', STR, 'false')]))
            links.append(f'  <hasGeneticAssociation rdf:resource="{iri}"/>')
        done.append(f'{len(alleles)} HLA association')

    if not done:
        continue
    open_k, close_k = span[ari]
    anchor = open_k + 1 if lines[open_k + 1].lstrip().startswith('<rdf:type ') else open_k
    after[anchor] = links
    note = f'Imported from {SOURCE} from {CURATOR}, ALPHABETICAL row {row}: {", ".join(done)}.'
    before[close_k] = values + [f'  <ARI_ChangeLog{STR}>{now} | {CURATOR} | {escape(note)}</ARI_ChangeLog>']
    for d in done:
        key = re.sub(r'^\d+ |\(.*\)| [A-Z]\w+$', '', d).strip()
        totals[key] = totals.get(key, 0) + 1

out = []
for k, line in enumerate(lines):
    out += before.get(k, [])
    out.append(line)
    out += after.get(k, [])
    if line == '<owl:ObjectProperty rdf:about="#affectsCellTissueOrOrgan"/>':
        out += ['', f'<owl:ObjectProperty rdf:about="#{ASSOC_PROP}"/>']
    if line == '<owl:AnnotationProperty rdf:about="#ARI_RefLink"/>':
        out += ['', f'<owl:AnnotationProperty rdf:about="#{ASSOC_TEXT}"/>']
end = len(out) - 1 - out[::-1].index('</rdf:RDF>')
end -= out[end - 1] == ''          # keep the blank line(s) before </rdf:RDF> at the end
for block in blocks:
    out[end:end] = block + ['']
    end += len(block) + 1
result = nl.join(out)
ET.fromstring(result)
for decl in (f'#{ASSOC_PROP}', f'#{ASSOC_TEXT}'):
    if f'rdf:about="{decl}"' not in result:
        raise SystemExit(f'Could not declare {decl}')
open(OWL_PATH, 'w', encoding='utf-8', newline='').write(result)
print(f'{len(before)} diseases, {len(blocks)} new items')
for k, n in sorted(totals.items()):
    print(f'  {k}: {n} diseases')
