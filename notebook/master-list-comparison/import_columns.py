"""HISTORICAL - kept as the record of the 2026-10-08 import; it refuses to run.

Import six immunology columns of the master list into ari_t1d.owl (now `ari_disease.owl`), for every disease that
report 10 matched: Target Antigens, Target Cells, Antibodies, Self Reactive Cells, Molecular
Mimicry and Environmental Triggers. Each line of a cell becomes one item, linked the way the
editor reads that category. Items a disease already has are not added again.

It reads and writes ari_t1d.owl line by line in the layout the file had before the canonical
formatter (#154), with relative `#` IRIs. Since #154 every IRI is absolute, so its `#` lookups and
the declaration lines it anchors on match nothing, and the items it writes would carry relative
IRIs. Do not re-run it; add or change items in the editor instead.

Usage: python import_columns.py <master.xlsx> <ari_t1d.owl> <10_Master_List_Comparison.xlsx>
"""
import re, sys
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from xml.sax.saxutils import escape

import openpyxl

from sheet_items import items, norm, same_item

raise SystemExit('Historical: written for the pre-#154 layout of ari_t1d.owl. '
                 'See the docstring.')

MASTER, OWL_PATH, REPORT = sys.argv[1:4]
SOURCE = 'ARI Master List 2026-10-08'
CURATOR = 'Leon Islas'
STR = ' rdf:datatype="http://www.w3.org/2001/XMLSchema#string"'
DISEASE_IRI = 'https://diseases.autoimmuneregistry.org/disease/ARI_'
A = '{http://www.aurint.org/ontologies/ari-disease#}'
RDF = '{http://www.w3.org/1999/02/22-rdf-syntax-ns#}'
RDFS = '{http://www.w3.org/2000/01/rdf-schema#}'
FREQ = re.compile(r'\s*(?:\(([^()]*\d[^()]*%[^()]*)\)|-\s*(\d[\d.,–-]*\s*%))\s*$')
T_CELL = re.compile(r'\bt\b|t cell|t-cell|t helper|\bth\d|treg|tfh|cd4|cd8|t lymph|intraepithelial|\bctl', re.I)

# What each column becomes. The new class and property for target cells follow the data
# model's affectsCellTissueOrOrgan / AffectedCellTissueOrOrgan.
#   kind: (link property, class, IRI prefix, data property for a frequency, (singular, plural))
KINDS = {
    'antigen': ('involvesAntigen', 'Antigen', 'Antigen', 'componentRelevance', ('target antigen', 'target antigens')),
    'cell': ('affectsCellTissueOrOrgan', 'AffectedCellTissueOrOrgan', 'Cell', 'componentRelevance',
             ('target cell', 'target cells')),
    'antibody': ('hasAntibody', 'Autoantibody', 'Ab', 'antibodyFrequency', ('antibody', 'antibodies')),
    'tcell': ('involvesTCell', 'TCellSubset', 'T', 'componentRelevance', ('self-reactive T cell', 'self-reactive T cells')),
    'apc': ('involvesAPC', 'APC', 'APC', 'componentRelevance',
            ('self-reactive B or other cell (as APC)', 'self-reactive B or other cells (as APC)')),
    'mimicry': ('hasEtiologyFactor', 'EtiologyOrigin', 'Et', None, ('molecular mimicry agent', 'molecular mimicry agents')),
    'trigger': ('hasEnvironmentalTrigger', 'EnvironmentalFactor', 'Env', 'likelihood',
                ('environmental trigger', 'environmental triggers')),
}
# Sheet header prefix -> function from a cell line to its kind.
COLUMNS = {
    'Target Antigens': lambda line: 'antigen',
    'Target Cells': lambda line: 'cell',
    'Antibodies': lambda line: 'antibody',
    'Self Reactive Cells': lambda line: 'tcell' if T_CELL.search(line) else 'apc',
    'Molecular Mimicry': lambda line: 'mimicry',
    'Enviromental Triggers': lambda line: 'trigger',
}

text = open(OWL_PATH, encoding='utf-8', newline='').read()
nl = '\r\n' if '\r\n' in text else '\n'
lines = text.split(nl)
root = ET.fromstring(text)
nodes = {e.get(RDF + 'about'): e for e in root if e.get(RDF + 'about')}
for decl in ('#affectsCellTissueOrOrgan', '#AffectedCellTissueOrOrgan'):
    if decl in nodes or f'rdf:about="{decl}"' in text:
        raise SystemExit(f'{decl} is already declared')


def existing(disease, kind):
    """What ARI already holds for the kind, as text comparable with a sheet line."""
    link = KINDS[kind][0]
    links = [link] if kind not in ('tcell', 'apc') else ['involvesTCell', 'involvesAPC']
    out = []
    for p in links:
        for r in disease.findall(A + p):
            n = nodes[r.get(RDF + 'resource')]
            if kind == 'mimicry':
                if n.findtext(RDFS + 'label') == 'Molecular mimicry':
                    out.append(n.findtext(A + 'etiologyDescription') or '')
            else:
                out.append(n.findtext(RDFS + 'label') or '')
    return out


ws = openpyxl.load_workbook(MASTER, data_only=True)['ALPHABETICAL']
headers = [str(c.value or '').strip() for c in ws[1]]
colidx = {}
for prefix in COLUMNS:
    hits = [k for k, h in enumerate(headers) if h.startswith(prefix)]
    if len(hits) != 1:
        raise SystemExit(f'Expected one column starting {prefix!r}, found {len(hits)}')
    colidx[prefix] = hits[0] + 1
match = [(r[0], r[2]) for r in openpyxl.load_workbook(REPORT)['Disease Match'].iter_rows(min_row=2, values_only=True)
         if r[0] and r[2] and r[4] != 'Duplicate row']

# Line index of each disease's opening tag and closing tag.
span = {}
for k, line in enumerate(lines):
    m = re.match(r'<([\w:]+) rdf:about="' + re.escape(DISEASE_IRI) + r'(\d{7})">$', line)
    if m:
        close = next(j for j in range(k + 1, len(lines)) if lines[j] == f'</{m.group(1)}>')
        span['ARI:' + m.group(2)] = (k, close)

now = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%S+00:00')
after, before, blocks, totals, skipped_total = {}, {}, [], {}, 0
for row, ari in match:
    disease = next(e for e in root if e.findtext(A + 'ARI_ID') == ari)
    counter, links, counts, skipped = {}, [], {}, 0
    for prefix, kind_of in COLUMNS.items():
        seen = set()
        for line in items(ws.cell(row, colidx[prefix]).value):
            kind = kind_of(line)
            link, cls, iri_prefix, freq_prop, _ = KINDS[kind]
            if norm(line) in seen:
                continue
            seen.add(norm(line))
            if any(same_item(line, x) for x in existing(disease, kind)):
                skipped += 1
                continue
            freq = FREQ.search(line) if freq_prop else None
            name = line[:freq.start()].strip() if freq else line
            counter[iri_prefix] = counter.get(iri_prefix, 0) + 1
            iri = f'#{iri_prefix}_{ari[4:]}_{counter[iri_prefix]}'
            if iri in nodes:
                raise SystemExit(f'{iri} already exists')
            block = [f'<owl:NamedIndividual rdf:about="{iri}">', f'  <rdf:type rdf:resource="#{cls}"/>']
            if kind == 'mimicry':
                block += [f'  <rdfs:label{STR}>Molecular mimicry</rdfs:label>',
                          f'  <etiologyOriginType{STR}>External</etiologyOriginType>',
                          f'  <etiologyDescription{STR}>{escape(name)}</etiologyDescription>']
            else:
                block.append(f'  <rdfs:label{STR}>{escape(name)}</rdfs:label>')
            if freq:
                block.append(f'  <{freq_prop}{STR}>{escape((freq.group(1) or freq.group(2)).strip())}</{freq_prop}>')
            block += [f'  <ARI_Obsolete{STR}>false</ARI_Obsolete>', '</owl:NamedIndividual>']
            blocks.append(block)
            links.append(f'  <{link} rdf:resource="{iri}"/>')
            counts[kind] = counts.get(kind, 0) + 1
    if not links:
        continue
    open_k, close_k = span[ari]
    anchor = open_k + 1 if lines[open_k + 1].lstrip().startswith('<rdf:type ') else open_k
    after[anchor] = links
    parts = [f'{n} {KINDS[k][4][n != 1]}' for k, n in counts.items()]
    note = f'Imported from {SOURCE} from {CURATOR}, ALPHABETICAL row {row}: {", ".join(parts)}.'
    if skipped:
        note += f' {skipped} more already in ARI were not added again.'
    before[close_k] = [f'  <ARI_ChangeLog{STR}>{now} | {CURATOR} | {escape(note)}</ARI_ChangeLog>']
    for k, n in counts.items():
        totals[k] = totals.get(k, 0) + n
    skipped_total += skipped

out = []
for k, line in enumerate(lines):
    out += before.get(k, [])
    out.append(line)
    out += after.get(k, [])
    if line == '<owl:ObjectProperty rdf:about="#involvesAntigen"/>':
        out += ['', '<owl:ObjectProperty rdf:about="#affectsCellTissueOrOrgan"/>']
    if line == '</owl:Class>' and lines[k - 2] == '<owl:Class rdf:about="#Antigen">':
        out += ['', '<owl:Class rdf:about="#AffectedCellTissueOrOrgan">',
                '  <rdfs:subClassOf rdf:resource="http://www.w3.org/2002/07/owl#Thing"/>', '</owl:Class>']
end = len(out) - 1 - out[::-1].index('</rdf:RDF>')
end -= out[end - 1] == ''          # keep the blank line(s) before </rdf:RDF> at the end
for block in blocks:
    out[end:end] = block + ['']
    end += len(block) + 1
result = nl.join(out)
ET.fromstring(result)
for decl in ('#affectsCellTissueOrOrgan', '#AffectedCellTissueOrOrgan'):
    if f'rdf:about="{decl}"' not in result:
        raise SystemExit(f'Could not declare {decl}')
open(OWL_PATH, 'w', encoding='utf-8', newline='').write(result)
print(f'{len(before)} diseases, {len(blocks)} items, {skipped_total} already in ARI')
for k, n in totals.items():
    print(f'  {KINDS[k][4][1]}: {n}')
