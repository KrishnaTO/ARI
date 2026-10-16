"""Compare what ARI publishes per disease with the fields of the immunological data model, and
say for each disease and field what is still to be researched.

Published   - ontologies/ari_t1d.owl, the record the editor serves
Data model  - data_model/Immuno-data-model-v4.xlsx (the field list below follows it row by row)
Not yet imported, but already in the repository:
  master    - ALPHABETICAL sheet of data/1-master/ARI Master List - 2026-10-08.xlsx, matched to ARI
              IDs through the Disease Match sheet of report 10
  report 11 - direct parent/child links of data/4-reports/11_Disease_Hierarchy.xlsx, less those
              rejected in notebook/disease-hierarchy/proposed_parents.tsv

Usage: python build_research_gaps.py
Writes data/4-reports/12_Research_Gaps.xlsx.
"""
import csv, os, re
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
ONTOLOGY = os.path.join(ROOT, "ontologies/ari_t1d.owl")
MASTER = os.path.join(ROOT, "data/1-master/ARI Master List - 2026-10-08.xlsx")
MATCH = os.path.join(ROOT, "data/4-reports/10_Master_List_Comparison.xlsx")
HIERARCHY = os.path.join(ROOT, "data/4-reports/11_Disease_Hierarchy.xlsx")
PROPOSED_PARENTS = os.path.join(ROOT, "notebook/disease-hierarchy/proposed_parents.tsv")
OUT = os.path.join(ROOT, "data/4-reports/12_Research_Gaps.xlsx")

A = "{http://www.aurint.org/ontologies/ari-disease#}"
RDF = "{http://www.w3.org/1999/02/22-rdf-syntax-ns#}"
RDFS = "{http://www.w3.org/2000/01/rdf-schema#}"

PUBLISHED, PARTIAL, IN_REPO, RESEARCH = "Published", "Partial", "In repo - import", "Research"
STATUSES = [PUBLISHED, PARTIAL, IN_REPO, RESEARCH]
COLOUR = {PUBLISHED: "C6EFCE", PARTIAL: "FFEB9C", IN_REPO: "BDD7EE", RESEARCH: "F4B6B6"}
MEANING = {
    PUBLISHED: "ARI publishes a value in the field the data model asks for.",
    PARTIAL: "ARI publishes something, but not what the data model asks for (see the note).",
    IN_REPO: "Not published, but the 2026-10-08 master list or report 11 already holds it. "
             "Needs importing and checking, not new research.",
    RESEARCH: "Nothing in ARI or anywhere in this repository. Needs new research.",
}

# ---- published record ----------------------------------------------------------------------
root = ET.parse(ONTOLOGY).getroot()
nodes = {}
for el in root:
    about = el.get(RDF + "about")
    if about:
        nodes[about.lstrip("#")] = el


def values(el, prop, ns=A):
    return [(c.get(RDF + "resource") or c.text or "").strip() for c in el.findall(ns + prop)]


def node_type(ref):
    return {t.get(RDF + "resource").lstrip("#") for t in nodes[ref.lstrip("#")].findall(RDF + "type")}


diseases = [el for el in root
            if any(t.get(RDF + "resource") == "#AutoimmuneDisease" for t in el.findall(RDF + "type"))
            and "true" not in values(el, "ARI_Obsolete")]
diseases.sort(key=lambda el: values(el, "ARI_ID")[0])

# ---- data already in the repository ----------------------------------------------------------
ws = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)["ALPHABETICAL"]
sheet = list(ws.iter_rows(values_only=True))
headers = [str(h or "").strip() for h in sheet[0]]


def col(prefix, nth=0):
    hits = [k for k, h in enumerate(headers) if h.startswith(prefix)]
    return hits[nth]


C_DEF = col("Autoimune Disease Definition (according to \nAutoimmune")
C_MESH = col("Mesh Terms")
C_VERDICT = col("Veredict")
C_PATHO = col("PATHOPHYSIOLOGY")
C_SYMPTOMS = col("Symptoms")
C_FM = col("Female")
C_AUTOINFLAM = col("Summary Evidence Autoinflammation")
C_GENES_AI = col("1. Genes/Genetic Regions", 1)
C_ASSOC = col("Association with other Autoimmune")
C_REGIONS = col("Countries and Regions")
C_LANDMARK = col("Landmark Studies")
EVIDENCE_GRADES = [k for k, h in enumerate(headers) if re.match(r"\d\. .*\n1\. Theoretical", h)]

match = {r[2]: r[0] for r in openpyxl.load_workbook(MATCH, read_only=True)["Disease Match"]
         .iter_rows(min_row=2, values_only=True) if r[0] and r[2] and r[4] != "Duplicate row"}
master = {ari: sheet[row - 1] for ari, row in match.items()}

rejected = {(r["Child ARI ID"], r["Parent ARI ID"])
            for r in csv.DictReader(open(PROPOSED_PARENTS, encoding="utf-8"), delimiter="\t")
            if r["Relation"] == "reject"}
parent_links = defaultdict(list)
rel = openpyxl.load_workbook(HIERARCHY, read_only=True)["Relationships"]
rel_rows = list(rel.iter_rows(values_only=True))
rh = list(rel_rows[0])
for r in rel_rows[1:]:
    # Direct links only: Indirect ones skip an intermediate disease, Umbrella ones point at the
    # catch-all "Autoimmune disease", and Reversed ones are outvoted by the opposite direction.
    if r[rh.index("Link")] == "Direct" and (r[rh.index("Child ARI ID")], r[rh.index("Parent ARI ID")]) not in rejected:
        review = f' - review: {r[rh.index("Review")]}' if r[rh.index("Review")] else ""
        parent_links[r[rh.index("Child ARI ID")]].append(
            f'{r[rh.index("Parent ARI ID")]} {r[rh.index("Parent")]} ({r[rh.index("Sources")]}){review}')

PLACEHOLDER = re.compile(r"^(none|unknown|unconfirmed|unknown / unconfirmed|n/?a|-|0(\.0)?)$", re.I)


def cell(ari, k):
    """The master list's text for this disease, or '' when it says nothing."""
    row = master.get(ari)
    text = str(row[k]).strip() if row and k < len(row) and row[k] is not None else ""
    lines = [l.strip().lstrip("•-*·").strip() for l in text.split("\n")]
    lines = [l for l in lines if l and not PLACEHOLDER.match(l)]
    return "\n".join(lines)


def short(text, n=90):
    text = " / ".join(text.split("\n"))
    return text if len(text) <= n else text[:n - 1] + "…"


# ---- field checks: each returns (status, note) ------------------------------------------------
def linked(el, prop, cls=None):
    refs = values(el, prop)
    return [r for r in refs if cls is None or cls in node_type(r)]


def by_link(prop, cls=None, repo_col=None, repo_what=None, note=""):
    def check(el, ari):
        n = len(linked(el, prop, cls))
        if n:
            return PUBLISHED, f"{n} item(s)"
        if repo_col is not None and cell(ari, repo_col):
            return IN_REPO, f"{repo_what}: {short(cell(ari, repo_col))}"
        return RESEARCH, note
    return check


def by_value(prop, repo_col=None, repo_what=None, note=""):
    def check(el, ari):
        vs = [v for v in values(el, prop) if v]
        if vs:
            return PUBLISHED, short("; ".join(vs), 60)
        if repo_col is not None and cell(ari, repo_col):
            return IN_REPO, f"{repo_what}: {short(cell(ari, repo_col))}"
        return RESEARCH, note
    return check


def nothing(note, repo=None):
    def check(el, ari):
        if repo:
            found = repo(ari)
            if found:
                return IN_REPO, found
        return RESEARCH, note
    return check


def parent(el, ari):
    if values(el, "hasParentDisease") or values(el, "hasParentCategory"):
        return PUBLISHED, short("; ".join(values(el, "hasParentDisease") + values(el, "hasParentCategory")), 60)
    if parent_links[ari]:
        return IN_REPO, "Report 11: " + "; ".join(parent_links[ari])
    return RESEARCH, "No parent in any local source - confirm it is top-level or find one"


QUALITY = {"High", "Moderate", "Low", "Expert Opinion"}


def evidence_quality(el, ari):
    vs = values(el, "evidenceQuality")
    if any(v in QUALITY for v in vs):
        return PUBLISHED, "; ".join(vs)
    grade = cell(ari, C_AUTOINFLAM)
    held = f"evidenceQuality holds the evidence type '{'; '.join(vs)}', not a grade. " if vs else ""
    if grade in ("Strong", "Moderate"):
        # Not imported: the column grades autoinflammation evidence, unconfirmed as the model's grade.
        return RESEARCH, held + f"Candidate: master 'Summary Evidence Autoinflammation' grades it '{grade}' - "                                "confirm the column is the evidence-quality grade"
    return RESEARCH, held + "Grade High / Moderate / Low / Expert Opinion"


CATEGORIES = {"Autoimmune", "Autoinflammatory", "Primary Immunodeficiency", "Immune Deficiency", "Dysregulation"}


def disease_type(el, ari):
    vs = values(el, "diseaseCategory")
    if any(v in CATEGORIES for v in vs):
        return PUBLISHED, "; ".join(vs)
    if vs:
        return PARTIAL, f"'{'; '.join(vs)}' - autoimmune status not yet confirmed"
    verdict = master[ari][C_VERDICT] if ari in master else None
    if verdict:
        return IN_REPO, f"Master verdict: {verdict}"
    return RESEARCH, "Autoimmune / Autoinflammatory / Primary Immunodeficiency"


def code(prop, label):
    def check(el, ari):
        vs = [v for v in values(el, prop) if v]
        if vs:
            return PUBLISHED, "; ".join(vs)
        return RESEARCH, f"No {label} mapping - map in the editor"
    return check


def symptoms(el, ari):
    n = len(linked(el, "hasSymptom"))
    if n:
        return PUBLISHED, f"{n} symptom(s)"
    if cell(ari, C_SYMPTOMS):
        return IN_REPO, f"Master Symptoms: {short(cell(ari, C_SYMPTOMS))}"
    return RESEARCH, ""


def objective_signs(el, ari):
    if cell(ari, C_SYMPTOMS):
        return RESEARCH, "No field yet. Master Symptoms mixes symptoms and signs - split them while reviewing"
    return RESEARCH, "No field yet"


def case_reports(el, ari):
    lm = cell(ari, C_LANDMARK)
    return RESEARCH, "No field yet" + (f". Starting point - master Landmark Studies: {short(lm, 70)}" if lm else "")


def prevalence(el, ari):
    # Each rate is an EpidemiologyProfile figure (measure, region, sex, rate, year).
    figures = [nodes[r.lstrip("#")] for r in linked(el, "hasEpidemiologyProfile", "EpidemiologyProfile")]
    rated = [f for f in figures if values(f, "epiRatePer100k")]
    if rated:
        return PUBLISHED, "; ".join(values(f, "label", RDFS)[0] for f in rated)
    desc = values(el, "ARI_PrevalenceDesc")
    if desc:
        return PARTIAL, f"Text only, no rate per 100,000: {short(desc[0], 70)}"
    return RESEARCH, "No prevalence or incidence"


def demographic_bias(el, ari):
    if values(el, "demographicBias"):
        return PUBLISHED, short(values(el, "demographicBias")[0], 60)
    found = []
    if cell(ari, C_FM):
        found.append(f"female:male ratio {cell(ari, C_FM)}")
    if cell(ari, C_REGIONS):
        found.append(f"regions: {short(cell(ari, C_REGIONS), 60)}")
    if found:
        return IN_REPO, "Master " + "; ".join(found)
    if values(el, "ageRange"):
        return PARTIAL, f"Only age of onset ({values(el, 'ageRange')[0]}); no sex or ancestry skew"
    return RESEARCH, "Sex, age and ancestry skew"


def mechanism(el, ari):
    n = len(linked(el, "hasPathwayStep"))
    if n:
        return PUBLISHED, f"{n} pathway step(s)"
    if cell(ari, C_PATHO):
        return IN_REPO, f"Master PATHOPHYSIOLOGY (numbered steps): {short(cell(ari, C_PATHO))}"
    return RESEARCH, "Mechanism with DOI reference"


def genetic(cls, what):
    def check(el, ari):
        n = len(linked(el, "hasGeneticAssociation", cls))
        if n:
            return PUBLISHED, f"{n} {what}"
        genes = [l for l in cell(ari, C_GENES_AI).split("\n") if l and (cls == "HLAAssociation") == bool(HLA.search(l))]
        if genes:
            return IN_REPO, f"Master genes column: {short(chr(10).join(genes))}"
        return RESEARCH, ""
    return check


HLA = re.compile(r"\bHLA\b|human leukocyte antigen", re.I)


def systemic(el, ari):
    tissues = values(el, "targetsTissue")
    if "#Tissue_Systemic" in tissues:
        return PARTIAL, "No field yet; targetsTissue is Systemic, so the answer is likely Yes"
    return RESEARCH, "No field yet (Yes / No / Variable)"


def in_vitro(what):
    def check(el, ari):
        graded = [headers[k].split("\n")[0] for k in EVIDENCE_GRADES if cell(ari, k)]
        note = "No field yet. The master's 1-10 evidence-grade columns are empty"
        if graded:
            note = f"No field yet. Master grades only: {', '.join(graded)}"
        return RESEARCH, f"{what} with DOI reference. {note}"
    return check


def associations(el, ari):
    links, names = linked(el, "associatedWithAutoimmune"), values(el, "ARI_AssociatedDisease")
    if names:
        return PUBLISHED, f"{len(names)} named, {len(links)} linked to ARI diseases"
    if cell(ari, C_ASSOC):
        return IN_REPO, f"Master: {short(cell(ari, C_ASSOC))}"
    return RESEARCH, ""


def definition(el, ari):
    if [v for v in values(el, "comment", RDFS) if v]:
        return PUBLISHED, ""
    if cell(ari, C_DEF):
        return IN_REPO, f"Master definition: {short(cell(ari, C_DEF))}"
    return RESEARCH, ""


def definition_source(el, ari):
    vs = values(el, "ARI_DefSource") + values(el, "ARI_Pubmed")
    if vs:
        return PUBLISHED, short(vs[0], 60)
    return RESEARCH, "Definition has no cited source" if values(el, "comment", RDFS) else ""


def synonyms(el, ari):
    if values(el, "ARI_Synonym"):
        return PUBLISHED, f"{len(values(el, 'ARI_Synonym'))} synonym(s)"
    if cell(ari, C_MESH):
        # Not imported: MeSH entry terms mix synonyms with sub- and broader types.
        return RESEARCH, f"Candidates for synonym review - master MeSH terms: {short(cell(ari, C_MESH))}"
    return RESEARCH, ""


# (group, data-model field, check). Group and field names follow Immuno-data-model-v4.xlsx;
# the first two rows are the core annotations of data/instructions.md.
FIELDS = [
    ("Core record", "Definition", definition),
    ("Core record", "Definition source", definition_source),
    ("Disease", "Synonyms", synonyms),
    ("Disease", "Parent", parent),
    ("Disease", "Evidence Quality", evidence_quality),
    ("Disease", "Genetic / Inherited", nothing("No field yet (Yes / No / Partial)")),
    ("Disease", "Disease Type", disease_type),
    ("Disease", "Association with Other Autoimmune Diseases", associations),
    ("Standard Codes", "SNOMED", code("ARI_SNOMED", "SNOMED")),
    ("Standard Codes", "DOID", code("ARI_DOID", "DOID")),
    ("Standard Codes", "UMLS", code("ARI_UMLS", "UMLS")),
    ("Standard Codes", "MONDO", code("ARI_MONDO", "MONDO")),
    ("Evidence", "In Vitro Animal Cell Evidence", in_vitro("Animal-model findings")),
    ("Evidence", "In Vitro Human Cell Evidence", in_vitro("Human-cell findings")),
    ("Clinical Profile", "Cardinal Symptoms", symptoms),
    ("Clinical Profile", "Objective Signs", objective_signs),
    ("Clinical Profile", "Case Reports", case_reports),
    ("Clinical Profile", "Clinical Biomarkers",
     by_link("hasBiomarker", note="Autoantibodies are listed under Antibodies; lab, biopsy, imaging and physiological tests are missing")),
    ("Epidemiology", "Prevalence", prevalence),
    ("Epidemiology", "Demographic Bias", demographic_bias),
    ("Pathogenesis", "Hypersensitivity Type", nothing("No field yet (Gell-Coombs Type I-IV)")),
    ("Pathogenesis", "Molecular Mechanism", mechanism),
    ("Etiology", "External / Environmental Trigger", by_link("hasEnvironmentalTrigger")),
    ("Etiology", "Inherited Disease Susceptibility",
     nothing("No field yet (sibling relative risk, twin concordance, familial aggregation)")),
    ("Etiology", "HLA Gene Region", genetic("HLAAssociation", "HLA association(s)")),
    ("Etiology", "Gene Region / Polymorphisms", genetic("GeneVariant", "gene variant(s)")),
    ("Innate Immune Component", "Innate Receptors", by_link("involvesInnateComponent")),
    ("Innate Immune Component", "Complement", by_link("involvesComplement")),
    ("Innate Immune Component", "Acute Phase Reactants", by_link("involvesAcutePhaseReactant")),
    ("Innate Immune Component", "NETosis", by_link("involvesNETosis")),
    ("Innate Immune Component", "Inflammasome", by_link("involvesInflammasome")),
    ("Adaptive Immunity", "T-Cell", by_link("involvesTCell")),
    ("Adaptive Immunity", "Antigen Presenting Cells", by_link("involvesAPC")),
    ("Adaptive Immunity", "Antibodies", by_link("hasAntibody")),
    ("Adaptive Immunity", "Antigens", by_link("involvesAntigen")),
    ("Signaling & Molecular", "Cytokines", by_link("involvesCytokine")),
    ("Signaling & Molecular", "Transcription Factors", by_link("involvesTranscriptionFactor")),
    ("Signaling & Molecular", "Receptors", by_link("involvesReceptor")),
    ("Pathophysiology", "Affected Cell, Tissue or Organ", by_link("affectsCellTissueOrOrgan")),
    ("Pathophysiology", "Systemic Involvement", systemic),
]

# Item links whose items the data model asks to carry a DOI / PMID reference.
CITED = ["hasSymptom", "hasEnvironmentalTrigger", "hasEtiologyFactor", "hasAntibody", "involvesTCell",
         "involvesAPC", "involvesAntigen", "affectsCellTissueOrOrgan"]


def citations(el):
    refs = [r for p in CITED for r in values(el, p)]
    cited = sum(1 for r in refs if nodes[r.lstrip("#")].find(A + "sourcePMID") is not None)
    return len(refs), cited


# ---- evaluate ---------------------------------------------------------------------------------
results = []      # (ari, label, category, [(status, note)], items, cited)
for el in diseases:
    ari = values(el, "ARI_ID")[0]
    row = [check(el, ari) for _, _, check in FIELDS]
    items, cited = citations(el)
    results.append((ari, el.findtext(RDFS + "label"), "; ".join(values(el, "diseaseCategory")), row, items, cited))

# ---- write ------------------------------------------------------------------------------------
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
FILL = {s: PatternFill("solid", fgColor=c) for s, c in COLOUR.items()}
WRAP = Alignment(wrap_text=True, vertical="top")
wb = openpyxl.Workbook()


def header(ws, names, widths):
    ws.append(names)
    for k, w in enumerate(widths, 1):
        ws.cell(1, k).font = BOLD
        ws.cell(1, k).fill = HEAD_FILL
        ws.cell(1, k).alignment = WRAP
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "A2"


# Summary: per field, how many diseases are in each state.
ws = wb.active
ws.title = "Summary"
header(ws, ["Group", "Data-model field"] + STATUSES + ["Still to do"], [24, 40, 11, 11, 15, 11, 12])
for k, (group, field, _) in enumerate(FIELDS):
    c = Counter(r[3][k][0] for r in results)
    ws.append([group, field] + [c[s] for s in STATUSES] + [len(results) - c[PUBLISHED]])
    for j, s in enumerate(STATUSES, 3):
        if c[s]:
            ws.cell(ws.max_row, j).fill = FILL[s]
ws.append([])
total_items = sum(r[4] for r in results)
total_cited = sum(r[5] for r in results)
ws.append(["Citations", f"{total_cited} of {total_items} published symptom, trigger, mimicry, antibody, cell and "
           f"antigen items carry a sourcePMID. The data model asks for a DOI reference on each."])
ws.append(["Diseases", f"{len(results)} published diseases (ARI_Obsolete excluded). "
           f"{len(master)} matched to a 2026-10-08 master-list row."])
ws.cell(ws.max_row - 1, 1).font = BOLD
ws.cell(ws.max_row, 1).font = BOLD

# Gap matrix: one row per disease, one coloured cell per field.
ws = wb.create_sheet("Gap Matrix")
header(ws, ["ARI ID", "Disease", "Disease Type"] + [f for _, f, _ in FIELDS] +
       ["Research", "In repo - import", "Items cited"],
       [12, 38, 14] + [6] * len(FIELDS) + [10, 10, 10])
for k in range(4, 4 + len(FIELDS)):
    ws.cell(1, k).alignment = Alignment(text_rotation=90, vertical="bottom", horizontal="center")
ws.row_dimensions[1].height = 190
ABBR = {PUBLISHED: "✓", PARTIAL: "~", IN_REPO: "M", RESEARCH: "R"}
for ari, label, cat, row, items, cited in results:
    c = Counter(s for s, _ in row)
    ws.append([ari, label, cat] + [ABBR[s] for s, _ in row] + [c[RESEARCH], c[IN_REPO], f"{cited}/{items}"])
    for k, (s, note) in enumerate(row, 4):
        cl = ws.cell(ws.max_row, k)
        cl.fill = FILL[s]
        cl.alignment = Alignment(horizontal="center")
ws.auto_filter.ref = ws.dimensions
ws.freeze_panes = "D2"

# To-do: every field still to be done, one line per disease and field.
ws = wb.create_sheet("To Do")
header(ws, ["ARI ID", "Disease", "Group", "Data-model field", "Status", "Note"], [12, 38, 22, 34, 15, 110])
for ari, label, _, row, _, _ in results:
    for (group, field, _), (s, note) in zip(FIELDS, row):
        if s != PUBLISHED:
            ws.append([ari, label, group, field, s, note])
            ws.cell(ws.max_row, 5).fill = FILL[s]
ws.auto_filter.ref = ws.dimensions

# Legend.
ws = wb.create_sheet("Legend")
header(ws, ["Status", "Matrix mark", "Meaning"], [16, 12, 110])
for s in STATUSES:
    ws.append([s, ABBR[s], MEANING[s]])
    ws.cell(ws.max_row, 1).fill = FILL[s]
ws.append([])
for line in [
    "Fields follow data_model/Immuno-data-model-v4.xlsx; Definition and its source are the core annotations of data/instructions.md.",
    "'No field yet' means ari_t1d.owl declares no property for the data-model field, so a value cannot be published until one is added.",
    "Gene Region and the two Gene Polymorphisms rows of the data model are one column here: ARI stores all three as GeneVariant.",
    "Published values are checked for presence only, not correctness. Evidence Quality and Disease Type are also checked against the data model's allowed values.",
    "Sources: ontologies/ari_t1d.owl; data/1-master/ARI Master List - 2026-10-08.xlsx (ALPHABETICAL) via report 10's Disease Match; report 11 Relationships. No online sources.",
]:
    ws.append([line])

wb.save(OUT)
print(f"{len(results)} diseases, {len(FIELDS)} fields -> {os.path.relpath(OUT, ROOT)}")
for (group, field, _), k in zip(FIELDS, range(len(FIELDS))):
    c = Counter(r[3][k][0] for r in results)
    print(f"  {field:44s} " + "  ".join(f"{s[:4]}={c[s]:3d}" for s in STATUSES))
print(f"  citations: {total_cited}/{total_items} items")
