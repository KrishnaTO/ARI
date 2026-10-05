"""Find parent/child relations between ARI core diseases.

Evidence sources (all local):
  master   - Diseases + Subtypes sheets of the ARI master list: a subtype/broad-class row
             whose name is another core disease's preferred name or synonym
  ontology - existing hasParentDisease links in ontologies/ari_t1d.owl
  snomed   - SNOMED is-a, via OMOP CONCEPT_ANCESTOR between the diseases' OMOP ConceptIDs
  doid     - DOID is_a (doid.owl subClassOf) between the diseases' DOID matches (report 6)

Usage: python build_disease_hierarchy.py <path to data/2-databases>
Writes data/4-reports/10_Disease_Hierarchy.xlsx and disease_hierarchy.csv next to this script.
"""
import csv, os, re, sys
import xml.etree.ElementTree as ET
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

if len(sys.argv) != 2:
    raise SystemExit(__doc__)
DB = sys.argv[1]
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
MASTER = os.path.join(ROOT, "data/1-master/ARI Master List V 2.1 - 2026-06-04.xlsx")
CORE = os.path.join(ROOT, "data/4-reports/1_Core_ARI_Diseases.xlsx")
DOID_MATCHES = os.path.join(ROOT, "notebook/ari-grounding/doid_matches_all.csv")
ONTOLOGY = os.path.join(ROOT, "ontologies/ari_t1d.owl")
OUT_XLSX = os.path.join(ROOT, "data/4-reports/10_Disease_Hierarchy.xlsx")
OUT_CSV = os.path.join(HERE, "disease_hierarchy.csv")
SOURCES = ["master", "ontology", "snomed", "doid"]

def norm(s):
    return re.sub(r"\s+", " ", str(s).replace("''", "'").replace("’", "'")).strip().lower()

# ---- core diseases ----------------------------------------------------------
core = {}          # ari -> dict
name_index = defaultdict(set)
rows = list(openpyxl.load_workbook(CORE, read_only=True).worksheets[0].iter_rows(values_only=True))
hdr = rows[0]
for r in rows[1:]:
    d = dict(zip(hdr, r))
    ari = d["ARI ID"]
    syns = [s for s in (d["Synonyms"] or "").split("; ") if s]
    omop = [int(x) for x in str(d["OMOP ConceptID"] or "").split(", ") if x.strip().isdigit()]
    core[ari] = {"name": d["Preferred Name"], "region": d["Tissue Region"], "omop": omop}
    for n in [d["Preferred Name"]] + syns:
        name_index[norm(n)].add(ari)
if not core:
    raise SystemExit("no core diseases read from report 1")

def lookup(name):
    hits = name_index.get(norm(name), set())
    return next(iter(hits)) if len(hits) == 1 else None

# evidence[(parent, child)][source] = detail
evidence = defaultdict(dict)
def add(parent, child, source, detail):
    if parent and child and parent != child:
        evidence[(parent, child)][source] = detail

# ---- master list ------------------------------------------------------------
wb = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
drows = list(wb["Diseases"].iter_rows(values_only=True))
dh = drows[0]
SUBTYPE = {"T", "TW", "TAW", "W", "WA"}
for r in drows[1:]:
    d = dict(zip(dh, r))
    syn, parent_name, name = d["syn"], d["Parent"], d["Disease"]
    if not parent_name or not name:
        continue
    group = lookup(parent_name)
    other = lookup(name)
    if syn in SUBTYPE:     # row names a subtype of its Parent group
        add(group, other, "master", f"Diseases sheet: '{name}' listed as subtype (syn={syn}) of '{parent_name}'")
    elif syn == "NN":      # row names an overbroad class that has this group as a subtype
        add(other, group, "master", f"Diseases sheet: '{name}' listed as broad class (syn=NN) of '{parent_name}'")
srows = list(wb["Subtypes"].iter_rows(values_only=True))
for r in srows[1:]:
    if r[1] and r[2]:
        add(lookup(r[1]), lookup(r[2]), "master", f"Subtypes sheet: '{r[2]}' is a subtype of '{norm(r[1])}'")

# ---- existing ontology links ------------------------------------------------
RDF = "{http://www.w3.org/1999/02/22-rdf-syntax-ns#}"
onto_names = {}
for ind in ET.parse(ONTOLOGY).getroot():
    about = ind.get(RDF + "about", "")
    if "/disease/ARI_" not in about:
        continue
    child = "ARI:" + about.rsplit("_", 1)[1]
    for e in ind:
        if e.tag.endswith("label"):
            onto_names[child] = e.text
    for e in ind:
        if e.tag.endswith("hasParentDisease"):
            parent = "ARI:" + e.get(RDF + "resource").rsplit("_", 1)[1]
            add(parent, child, "ontology", "hasParentDisease in ontologies/ari_t1d.owl")

# ---- SNOMED via OMOP CONCEPT_ANCESTOR --------------------------------------
omop_to_ari = defaultdict(set)
for ari, d in core.items():
    for c in d["omop"]:
        omop_to_ari[c].add(ari)
omop_names = {}
primary_snomed = set()   # pairs linked through both diseases' first-listed OMOP ConceptID
with open(os.path.join(DB, "snomed/CONCEPT.csv"), encoding="utf-8") as f:
    for r in csv.DictReader(f, delimiter="\t", quoting=csv.QUOTE_NONE):
        cid = int(r["concept_id"])
        if cid in omop_to_ari:
            omop_names[cid] = (r["concept_name"], r["concept_code"])
with open(os.path.join(DB, "snomed/CONCEPT_ANCESTOR.csv"), encoding="utf-8") as f:
    next(f)
    for line in f:
        a, d_, lo, _ = line.rstrip("\n").split("\t")
        a, d_ = int(a), int(d_)
        if a in omop_to_ari and d_ in omop_to_ari and a != d_:
            for pa in omop_to_ari[a]:
                for ch in omop_to_ari[d_]:
                    if a == core[pa]["omop"][0] and d_ == core[ch]["omop"][0]:
                        primary_snomed.add((pa, ch))
                    add(pa, ch, "snomed",
                        f"SNOMED {omop_names[d_][1]} '{omop_names[d_][0]}' is-a {omop_names[a][1]} "
                        f"'{omop_names[a][0]}' ({lo} level(s))")

# ---- DOID is_a --------------------------------------------------------------
doid_to_ari = defaultdict(set)
with open(DOID_MATCHES, encoding="utf-8") as f:
    for r in csv.DictReader(f):
        if r["DOID"] and r["ARI ID"] in core:
            doid_to_ari[r["DOID"]].add(r["ARI ID"])
OWL = "{http://www.w3.org/2002/07/owl#}"
RDFS = "{http://www.w3.org/2000/01/rdf-schema#}"
doid_parents, doid_labels = defaultdict(set), {}
for cls in ET.parse(os.path.join(DB, "doid.owl")).getroot().findall(OWL + "Class"):
    about = cls.get(RDF + "about", "")
    if "/DOID_" not in about:
        continue
    did = "DOID:" + about.rsplit("DOID_", 1)[1]
    lab = cls.find(RDFS + "label")
    doid_labels[did] = lab.text if lab is not None else ""
    for sc in cls.findall(RDFS + "subClassOf"):
        res = sc.get(RDF + "resource", "")
        if "/DOID_" in res:
            doid_parents[did].add("DOID:" + res.rsplit("DOID_", 1)[1])

def doid_ancestors(did):
    seen, frontier, depth = {}, {did}, 0
    while frontier:
        depth += 1
        nxt = set()
        for x in frontier:
            for p in doid_parents[x]:
                if p not in seen:
                    seen[p] = depth
                    nxt.add(p)
        frontier = nxt
    return seen

for did, children in doid_to_ari.items():
    for anc, lvl in doid_ancestors(did).items():
        for pa in doid_to_ari.get(anc, ()):
            for ch in children:
                add(pa, ch, "doid", f"DOID {did} '{doid_labels[did]}' is_a {anc} '{doid_labels[anc]}' ({lvl} level(s))")

# ---- review flags, conflicts and direct links ------------------------------
# A link supported only by SNOMED through an additional (not first-listed) OMOP code is
# flagged: the extra code is often a related concept rather than the disease itself.
additional_only = {p for p, ev in evidence.items() if set(ev) == {"snomed"} and p not in primary_snomed}
# "Autoimmune disease" is the catch-all class; its links are kept but left out of the tree.
umbrella = {a for a, d in core.items() if norm(d["name"]) == "autoimmune disease"}
two_way = {p for p in evidence if (p[1], p[0]) in evidence}
superseded = {p for p in two_way if p in additional_only and (p[1], p[0]) not in additional_only}
conflicts = {p for p in two_way if p not in superseded and (p[1], p[0]) not in superseded}
edges = {p for p in evidence if p not in conflicts and p not in superseded and p[0] not in umbrella}
parents_of = defaultdict(set)
for pa, ch in edges:
    parents_of[ch].add(pa)

def all_ancestors(x):
    seen, stack = set(), list(parents_of[x])
    while stack:
        y = stack.pop()
        if y not in seen:
            seen.add(y)
            stack.extend(parents_of[y])
    return seen

def via(pa, ch):
    """Intermediate listed diseases between pa and ch, if the link is indirect."""
    return sorted(m for m in parents_of[ch] if m != pa and pa in all_ancestors(m) | {m} and m != pa)

def label(a):
    return core[a]["name"] if a in core else onto_names.get(a, a) + " (ontology only)"

records = []
for (pa, ch), ev in sorted(evidence.items(), key=lambda kv: (label(kv[0][0]).lower(), label(kv[0][1]).lower())):
    if pa in umbrella:
        status = "Umbrella"
    elif (pa, ch) in superseded:
        status = "Reversed"
    elif (pa, ch) in conflicts:
        status = "Conflict"
    else:
        status = "Indirect" if via(pa, ch) else "Direct"
    records.append({
        "Parent ARI ID": pa, "Parent": label(pa), "Child ARI ID": ch, "Child": label(ch),
        "Link": status, "Via": "; ".join(f"{m} {label(m)}" for m in via(pa, ch)) if status == "Indirect" else "",
        "Sources": ", ".join(s for s in SOURCES if s in ev), "Source count": len(ev),
        "Review": "SNOMED link only through an additional OMOP code" if (pa, ch) in additional_only else "",
        **{s: ev.get(s, "") for s in SOURCES},
    })

with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(records[0]))
    w.writeheader()
    w.writerows(records)

# ---- workbook ---------------------------------------------------------------
FILL = {"master": "C6EFCE", "ontology": "BDD7EE", "snomed": "FFE699", "doid": "F8CBAD"}
LINK_FILL = {"Direct": "FFFFFF", "Indirect": "EDEDED", "Conflict": "FF9999", "Reversed": "FF9999", "Umbrella": "EDEDED"}
bold = Font(bold=True)
wb = openpyxl.Workbook()

ws = wb.active
ws.title = "Hierarchy"
ws.append(["Disease tree (direct links only; a disease with several parents appears under each)"])
ws["A1"].font = bold
ws.append(["Level", "ARI ID", "Disease", "Tissue Region", "Sources (link to parent)", "Review"])
for c in ws[2]:
    c.font = bold
direct = [r for r in records if r["Link"] == "Direct"]
kids = defaultdict(list)
src_of = {}
for r in direct:
    kids[r["Parent ARI ID"]].append(r["Child ARI ID"])
    src_of[(r["Parent ARI ID"], r["Child ARI ID"])] = (r["Sources"], r["Review"])
child_set = {r["Child ARI ID"] for r in direct}
roots = sorted({r["Parent ARI ID"] for r in direct} - child_set, key=lambda a: label(a).lower())

def walk(a, level, parent):
    region = core[a]["region"] if a in core else ""
    srcs, review = src_of.get((parent, a), ("", ""))
    ws.append([level, a, "    " * level + label(a), region, srcs, review])
    row = ws.max_row
    if level == 0:
        for c in ws[row]:
            c.font = bold
    first = srcs.split(", ")[0] if srcs else ""
    if first:
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=FILL[first])
    for k in sorted(kids[a], key=lambda x: label(x).lower()):
        walk(k, level + 1, a)

for r in roots:
    walk(r, 0, None)
for col, wdt in zip("ABCDEF", [7, 13, 70, 22, 30, 45]):
    ws.column_dimensions[col].width = wdt
ws.freeze_panes = "A3"

rs = wb.create_sheet("Relationships")
cols = list(records[0])
rs.append(cols)
for c in rs[1]:
    c.font = bold
for r in records:
    rs.append([r[c] for c in cols])
    row = rs.max_row
    rs.cell(row, cols.index("Link") + 1).fill = PatternFill("solid", fgColor=LINK_FILL[r["Link"]])
    for s in SOURCES:
        if r[s]:
            rs.cell(row, cols.index(s) + 1).fill = PatternFill("solid", fgColor=FILL[s])
for i, wdt in enumerate([13, 40, 13, 40, 10, 40, 26, 8, 30, 60, 30, 70, 70], 1):
    rs.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
for row in rs.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)
rs.freeze_panes = "A2"
rs.auto_filter.ref = rs.dimensions

lg = wb.create_sheet("Legend")
for line in [
    ["Source", "Colour", "Meaning"],
    ["master", "", "ARI Master List V 2.1 (Diseases sheet syn codes T/TW/TAW/W/WA = subtype, NN = broad class; Subtypes sheet). Names matched exactly to a core disease's preferred name or synonym."],
    ["ontology", "", "Existing hasParentDisease link in ontologies/ari_t1d.owl. Diseases marked '(ontology only)' were added to the ontology after the master list."],
    ["snomed", "", "SNOMED CT is-a hierarchy between the diseases' OMOP ConceptIDs, from local data/2-databases/snomed/CONCEPT_ANCESTOR.csv (OMOP/Athena export)."],
    ["doid", "", "Human Disease Ontology is_a hierarchy between the diseases' DOID matches (report 6), from local data/2-databases/doid.owl (release 2026-04-30)."],
    [],
    ["Link", "", "Meaning"],
    ["Direct", "", "No other listed disease sits between parent and child."],
    ["Indirect", "", "Parent is also an ancestor of an intermediate listed disease (see Via); omitted from the tree."],
    ["Conflict", "", "Sources disagree on direction (both A>B and B>A found); review manually; omitted from the tree."],
    ["Reversed", "", "Opposite direction also found and is better supported (this direction rests only on an additional SNOMED code); omitted from the tree."],
    ["Umbrella", "", "Parent is the catch-all 'Autoimmune disease'; kept for completeness, omitted from the tree."],
    [],
    ["Review", "", "'SNOMED link only through an additional OMOP code': the link relies on a second/third code listed for a disease, not its main code. Often that code names a related condition (e.g. Lofgren syndrome under Sarcoidosis), so check the code assignment in the master list."],
]:
    lg.append(line)
for i, s in enumerate(SOURCES, 2):
    lg.cell(i, 2).fill = PatternFill("solid", fgColor=FILL[s])
for i, s in enumerate(["Direct", "Indirect", "Conflict", "Reversed", "Umbrella"], 8):
    lg.cell(i, 2).fill = PatternFill("solid", fgColor=LINK_FILL[s])
lg["A1"].font = lg["A7"].font = bold
lg.column_dimensions["A"].width = 12
lg.column_dimensions["C"].width = 120
wb.save(OUT_XLSX)

n = defaultdict(int)
for r in records:
    n[r["Link"]] += 1
    for s in SOURCES:
        n[s] += bool(r[s])
diseases = {r["Parent ARI ID"] for r in records} | {r["Child ARI ID"] for r in records}
print(f"links: {len(records)} {dict(n)}; diseases involved: {len(diseases)}; tree roots: {len(roots)}")
