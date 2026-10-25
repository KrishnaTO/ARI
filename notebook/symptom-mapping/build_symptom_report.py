"""Write the Symptoms and Symptom Mapping sheets of report 4 from the ontology (ARI#132).

The OWL is the record of every symptom mapping decision; this script only reports it:
  - ontologies/ari_symptoms.owl: one ARIS class per symptom concept. Its rdfs:label is the symptom,
    skos:exactMatch / skos:broadMatch name its SYMP and HPO terms (no HPO match = "no term"),
    skos:altLabel holds the strings merged into it, and rdfs:subClassOf its parents: other ARIS
    classes, or SYMP terms when it has no ARIS parent.
  - ontologies/ari_symp.owl: labels of the SYMP terms those classes name.
  - ontologies/ari_disease.owl: the Symptom individuals, their labels and ARI_NoConcept
    ("not a symptom"), and the diseases that list them with hasSymptom.
  - hp.obo: HPO labels, since the OWL references HPO terms by IRI only.
  - symp.owl: the order of a SYMP term's own parents, which ROBOT's layout does not keep, so a
    class with several SYMP parents lists them in the release's order.

Writes data/4-reports/4_Additional_Info_Index.xlsx:
  - Symptoms sheet: the mapping and parent columns after the existing ones, for each row whose
    symptom string is a class label or altLabel, or the label of a "not a symptom" individual.
    A string the ontology does not know keeps the cells it has.
  - Symptom Mapping sheet: one row per symptom string (class labels and altLabels, Symptom
    individual labels and Symptoms-sheet strings), with the number of diseases listing it in the
    Symptoms sheet and in the ontology. A string the ontology does not know keeps the mapping
    columns of its existing row.

Usage: python build_symptom_report.py <path to data/2-databases>
The database folder is not tracked in git, so pass the path of a checkout holding symp.owl
(http://purl.obolibrary.org/obo/symp.owl) and hp.obo (http://purl.obolibrary.org/obo/hp.obo).
"""
import collections, os, sys
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.styles import Font, PatternFill

from build_symp_import import obo_iri, read_hp, read_symp

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SYMPTOMS = os.path.join(ROOT, "ontologies", "ari_symptoms.owl")
SYMP_SLICE = os.path.join(ROOT, "ontologies", "ari_symp.owl")
DISEASE = os.path.join(ROOT, "ontologies", "ari_disease.owl")
REPORT = os.path.join(ROOT, "data", "4-reports", "4_Additional_Info_Index.xlsx")

ARIS = "http://www.aurint.org/ontologies/ari-symptoms#ARIS_"
ARI_BASE = "http://www.aurint.org/ontologies/ari-disease#"
OBO = "http://purl.obolibrary.org/obo/"
NOT_A_SYMPTOM = "not a symptom"
RDF = "{http://www.w3.org/1999/02/22-rdf-syntax-ns#}"
OWL = "{http://www.w3.org/2002/07/owl#}"
RDFS = "{http://www.w3.org/2000/01/rdf-schema#}"
SKOS = "{http://www.w3.org/2004/02/skos/core#}"
OBOINOWL = "{http://www.geneontology.org/formats/oboInOwl#}"
ARI = "{%s}" % ARI_BASE

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
SYMPTOM_COLUMNS = [("Symptom ID", 14), ("SYMP ID", 14), ("SYMP Label", 30), ("SYMP Mapping", 14),
                   ("HPO ID", 14), ("HPO Label", 30), ("HPO Mapping", 14),
                   ("Parent Symptom", 30), ("Parent ID", 16), ("Merged Into", 30)]
MAPPING_COLUMNS = [("Symptom", 50), ("Symptom ID", 14), ("Mapping", 14), ("SYMP ID", 14), ("SYMP Label", 30),
                   ("HPO ID", 14), ("HPO Label", 30), ("HPO Mapping", 14),
                   ("Parent Symptom", 30), ("Parent ID", 16), ("Merged Into", 30),
                   ("Diseases (Symptoms sheet)", 14), ("Diseases (ontology)", 14)]
LINKED = {"SYMP ID", "HPO ID", "Parent ID"}


def curie(iri):
    return iri[len(OBO):].replace("_", ":")


def read_symp_labels():
    """SYMP CURIE -> label, from the ari_symp.owl slice."""
    root = ET.parse(SYMP_SLICE).getroot()
    return {c.findtext(OBOINOWL + "id"): c.findtext(RDFS + "label") for c in root.findall(OWL + "Class")}


def match(cls, prefix):
    """(CURIE, "exact" | "broader") of the class's one match to a term under prefix, or None."""
    found = [(curie(m.get(RDF + "resource")), "exact" if m.tag == SKOS + "exactMatch" else "broader")
             for tag in (SKOS + "exactMatch", SKOS + "broadMatch") for m in cls.findall(tag)
             if m.get(RDF + "resource").startswith(prefix)]
    if len(found) > 1:
        raise SystemExit(f"{cls.get(RDF + 'about')} has {len(found)} matches to {prefix}")
    return found[0] if found else None


def concepts(symptoms_root, symp_labels, hp_labels, symp_order):
    """Symptom string -> the values of its row: every ARIS class label and altLabel.

    symp_order: SYMP CURIE -> its parents in release order, to order a class's SYMP parents.
    """
    classes = {c.get(RDF + "about"): c for c in symptoms_root.findall(OWL + "Class")
               if c.get(RDF + "about").startswith(ARIS)}
    label = {iri: c.findtext(RDFS + "label") for iri, c in classes.items()}
    ident = {iri: c.findtext(OBOINOWL + "id") for iri, c in classes.items()}
    rows = {}

    def add(name, values):
        if name in rows:
            raise SystemExit(f"{name!r} is the label or altLabel of more than one ARIS class")
        rows[name] = values

    for iri, cls in classes.items():
        symp = match(cls, OBO + "SYMP_")
        if symp is None:
            raise SystemExit(f"{iri} has no SYMP match")
        hp = match(cls, OBO + "HP_")
        targets = [p.get(RDF + "resource") for p in cls.findall(RDFS + "subClassOf")]
        parents = [(label[t], ident[t]) for t in targets if t in classes]
        if not parents:
            symp_parents = [curie(t) for t in targets if t.startswith(OBO + "SYMP_")]
            if len(symp_parents) > 1:
                order = symp_order[symp[0]][1]
                if not set(symp_parents) <= set(order):
                    raise SystemExit(f"{iri}: SYMP parents {symp_parents} are not all parents of {symp[0]}")
                symp_parents.sort(key=order.index)
            parents = [(symp_labels[p], p) for p in symp_parents]
        values = {"Symptom ID": ident[iri], "Mapping": symp[1], "SYMP ID": symp[0], "SYMP Label": symp_labels[symp[0]],
                  "HPO ID": hp[0] if hp else "", "HPO Label": hp_labels[hp[0]] if hp else "",
                  "HPO Mapping": hp[1] if hp else "no term",
                  "Parent Symptom": "; ".join(p[0] for p in parents),
                  "Parent ID": "; ".join(p[1] for p in parents), "Merged Into": ""}
        add(label[iri], values)
        for alt in cls.findall(SKOS + "altLabel"):
            add(alt.text, dict(values, **{"Merged Into": label[iri]}))
    return rows


def read_individuals():
    """Symptom individual labels, those marked "not a symptom", and label -> ARI IDs of diseases listing it."""
    root = ET.parse(DISEASE).getroot()
    labels, not_symptoms = {}, set()
    for ind in root.findall(OWL + "NamedIndividual"):
        if any(t.get(RDF + "resource") == ARI_BASE + "Symptom" for t in ind.findall(RDF + "type")):
            name = ind.findtext(RDFS + "label").strip()
            labels[ind.get(RDF + "about")] = name
            if ind.findtext(ARI + "ARI_NoConcept") == NOT_A_SYMPTOM:
                not_symptoms.add(name)
    diseases = collections.defaultdict(set)
    for ind in root:
        for ref in ind.findall(ARI + "hasSymptom"):
            diseases[labels[ref.get(RDF + "resource")]].add(ind.findtext(ARI + "ARI_ID"))
    return set(labels.values()), not_symptoms, diseases


def header(ws, col, text, width):
    cell = ws.cell(row=1, column=col, value=text)
    cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.column_dimensions[cell.column_letter].width = width


def link(cell, curie):
    if curie.startswith(("SYMP:", "HP:")):
        cell.hyperlink, cell.font = obo_iri(curie), LINK_FONT


def sheet_values(rows, name):
    """The SYMPTOM_COLUMNS values of a known string: a concept row, or "not a symptom" alone."""
    if name in rows:
        r = rows[name]
        return [r["Symptom ID"], r["SYMP ID"], r["SYMP Label"], r["Mapping"], r["HPO ID"], r["HPO Label"],
                r["HPO Mapping"], r["Parent Symptom"], r["Parent ID"], r["Merged Into"]]
    return ["", "", "", NOT_A_SYMPTOM, "", "", "", "", "", ""]


def write_symptoms_sheet(ws, rows, not_symptoms):
    """Fill the mapping columns of the Symptoms sheet; return symptom string -> ARI IDs listing it there."""
    first = 8  # after ARI ID, seq, disease, symptom, latin, col4, col5
    if [c.value for c in ws[1]][3] != "symptom":
        raise SystemExit("Symptoms sheet layout changed: column D is no longer 'symptom'")
    for i, (text, width) in enumerate(SYMPTOM_COLUMNS):
        header(ws, first + i, text, width)
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=first + len(SYMPTOM_COLUMNS) - 1).column_letter}1"
    report_diseases = collections.defaultdict(set)
    for row in ws.iter_rows(min_row=2):
        name, kind = row[3].value, row[4].value
        if not name or kind == "cat":  # cat rows name a body system, not a symptom
            continue
        name = name.strip()
        report_diseases[name].add(row[0].value)
        if name not in rows and name not in not_symptoms:
            continue  # not in the ontology: the row keeps its cells
        for i, v in enumerate(sheet_values(rows, name)):
            cell = ws.cell(row=row[0].row, column=first + i, value=v or None)
            cell.hyperlink = None
            if SYMPTOM_COLUMNS[i][0] in LINKED:
                link(cell, v)
    return report_diseases


def write_mapping_sheet(wb, rows, individual_labels, not_symptoms, report_diseases, ontology_diseases):
    """Rewrite the Symptom Mapping sheet: one row per symptom string."""
    old = {r[0]: r[1:11] for r in wb["Symptom Mapping"].iter_rows(min_row=2, values_only=True)}
    names = set(rows) | individual_labels | set(report_diseases)
    del wb["Symptom Mapping"]
    sm = wb.create_sheet("Symptom Mapping", wb.sheetnames.index("Symptoms") + 1)
    for i, (text, width) in enumerate(MAPPING_COLUMNS, 1):
        header(sm, i, text, width)
    last = sm.cell(row=1, column=len(MAPPING_COLUMNS)).column_letter
    sm.freeze_panes, sm.auto_filter.ref = "A2", f"A1:{last}{len(names) + 1}"
    for n, name in enumerate(sorted(names, key=lambda s: (s.casefold(), s)), 2):
        if name in rows:
            r = rows[name]
            mapping = [r["Symptom ID"], r["Mapping"], r["SYMP ID"], r["SYMP Label"], r["HPO ID"], r["HPO Label"],
                       r["HPO Mapping"], r["Parent Symptom"], r["Parent ID"], ""]
            if r["Merged Into"]:  # a merged string's mapping is on the row it merges into
                mapping = ["", "merged"] + [""] * 7 + [r["Merged Into"]]
        elif name in not_symptoms:
            mapping = ["", NOT_A_SYMPTOM] + [""] * 8
        elif name in old:
            mapping = ["" if v is None else v for v in old[name]]
        else:
            raise SystemExit(f"{name!r} is in neither the ontology nor the Symptom Mapping sheet")
        values = [name] + mapping + [len(report_diseases[name]), len(ontology_diseases[name])]
        for i, v in enumerate(values, 1):
            cell = sm.cell(row=n, column=i, value=v if v != "" else None)
            if MAPPING_COLUMNS[i - 1][0] in LINKED:
                link(cell, v)
    return len(names)


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    symp_order, _ = read_symp(os.path.join(sys.argv[1], "symp.owl"))
    hp_labels, _ = read_hp(os.path.join(sys.argv[1], "hp.obo"))
    rows = concepts(ET.parse(SYMPTOMS).getroot(), read_symp_labels(), hp_labels, symp_order)
    individual_labels, not_symptoms, ontology_diseases = read_individuals()
    both = sorted(not_symptoms & set(rows))
    if both:
        raise SystemExit(f"Marked {NOT_A_SYMPTOM!r} yet a class label or altLabel: {both}")
    unknown = sorted(individual_labels - set(rows) - not_symptoms)
    if unknown:
        raise SystemExit(f"Symptom individuals with neither a concept nor {NOT_A_SYMPTOM!r}: {unknown}")

    wb = openpyxl.load_workbook(REPORT)
    report_diseases = write_symptoms_sheet(wb["Symptoms"], rows, not_symptoms)
    count = write_mapping_sheet(wb, rows, individual_labels, not_symptoms, report_diseases, ontology_diseases)
    wb.save(REPORT)
    sheet_only = sorted(set(report_diseases) - set(rows) - individual_labels)
    print(f"{len(rows)} class labels and altLabels, {len(not_symptoms)} not a symptom, "
          f"{len(sheet_only)} Symptoms-sheet strings the ontology does not know; "
          f"{count} rows in Symptom Mapping -> {os.path.relpath(REPORT, ROOT)}")


if __name__ == "__main__":
    main()
